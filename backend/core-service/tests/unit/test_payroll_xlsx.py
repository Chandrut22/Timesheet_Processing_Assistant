"""The payroll export as an Excel workbook: same rows as the CSV, real numbers and
dates, a totals row, and no formula injection."""

import uuid
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

import openpyxl

from src.core.services.payroll_csv import CSV_COLUMNS, ExportRow, build_csv
from src.core.services.payroll_xlsx import SHEET_TITLE, build_xlsx, xlsx_filename

WEEK = date(2026, 9, 14)


def row(code: str, name: str, regular: str, overtime: str = "0") -> ExportRow:
    return ExportRow(
        timesheet_id=uuid.uuid4(),
        employee_id=uuid.uuid4(),
        employee_code=code,
        employee_name=name,
        regular=Decimal(regular),
        overtime=Decimal(overtime),
        double_time=Decimal("0"),
        holiday=Decimal("0"),
        approved_by="Ravi Reviewer",
        approved_at=datetime(2026, 9, 21, 10, 30, tzinfo=UTC),
        currency="INR",
        hourly_rate=Decimal("500"),
        overtime_rate=Decimal("750"),
        double_time_rate=Decimal("1000"),
        holiday_rate=Decimal("1000"),
        regular_amount=Decimal(regular) * 500,
        overtime_amount=Decimal(overtime) * 750,
        double_time_amount=Decimal("0"),
        holiday_amount=Decimal("0"),
    )


def sheet_for(rows: list[ExportRow]) -> openpyxl.worksheet.worksheet.Worksheet:
    data = build_xlsx(build_csv("ABC Staffing", WEEK, rows))
    book = openpyxl.load_workbook(BytesIO(data))
    return book[SHEET_TITLE]


def col(name: str) -> int:
    return CSV_COLUMNS.index(name) + 1


def test_same_header_and_rows_as_the_csv() -> None:
    sheet = sheet_for(
        [row("E002", "Grace Okoye", "40", "2"), row("E001", "Ravi Kumar", "38")]
    )
    assert [c.value for c in sheet[1]] == CSV_COLUMNS
    assert sheet.cell(2, col("employee_code")).value == "E001"  # sorted like the CSV
    assert sheet.cell(3, col("employee_name")).value == "Grace Okoye"
    assert sheet.max_row == 4  # header + 2 rows + totals


def test_numbers_and_dates_are_real_values() -> None:
    sheet = sheet_for([row("E001", "Ravi Kumar", "40", "2.5")])
    overtime = sheet.cell(2, col("overtime_hours"))
    assert overtime.value == 2.5 and overtime.number_format == "0.00"
    gross = sheet.cell(2, col("gross_amount"))
    assert gross.value == 21875.0 and gross.number_format == "#,##0.00"
    assert sheet.cell(2, col("week_start")).value == datetime(2026, 9, 14)
    assert sheet.cell(2, col("approved_at")).value == datetime(2026, 9, 21, 10, 30)


def test_totals_row_adds_hours_and_amounts_not_rates() -> None:
    sheet = sheet_for([row("E001", "A", "40", "2"), row("E002", "B", "38")])
    total = sheet.max_row
    assert sheet.cell(total, 1).value == "Total"
    assert sheet.cell(total, col("regular_hours")).value == 78.0
    assert sheet.cell(total, col("total_hours")).value == 80.0
    assert sheet.cell(total, col("gross_amount")).value == 40500.0
    assert sheet.cell(total, col("hourly_rate")).value is None
    assert sheet.cell(total, 1).font.bold


def test_a_formula_looking_name_stays_text() -> None:
    sheet = sheet_for([row("E001", '=HYPERLINK("http://x")', "8")])
    cell = sheet.cell(2, col("employee_name"))
    assert cell.value == '=HYPERLINK("http://x")'  # shown as written, without the quote
    assert cell.data_type == "s"  # text, never a formula Excel would run


def test_header_frozen_with_filters() -> None:
    sheet = sheet_for([row("E001", "A", "8")])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:W2"  # data rows only, not the totals


def test_csv_with_byte_order_mark_and_old_exports() -> None:
    text = "﻿" + build_csv("ABC Staffing", WEEK, [row("E001", "A", "8")])
    book = openpyxl.load_workbook(BytesIO(build_xlsx(text)))
    assert book[SHEET_TITLE].cell(1, 1).value == "company"


def test_filename() -> None:
    assert xlsx_filename("payroll_abc_2026-09-14_1a2b3c4d.csv") == (
        "payroll_abc_2026-09-14_1a2b3c4d.xlsx"
    )
