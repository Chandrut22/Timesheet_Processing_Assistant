"""Large documents: page / sheet limits before reading, the document_too_large flag,
"too long for the model" errors, and cut-off AI answers."""

import asyncio
from io import BytesIO
from types import SimpleNamespace
from typing import Any, cast

import openpyxl
import pypdfium2
import pytest

from src.config.settings import Settings, settings
from src.core.services.failures import AnswerTooLongError, explain
from src.data.clients.document_limits import (
    DocumentTooLargeError,
    prepare,
    too_large_flag,
)


def pdf(pages: int) -> bytes:
    document = pypdfium2.PdfDocument.new()
    for _ in range(pages):
        document.new_page(595, 842)
    out = BytesIO()
    document.save(out)
    document.close()
    return out.getvalue()


def workbook(sheets: dict[str, list[list[Any]]], hidden: tuple[str, ...] = ()) -> bytes:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
        if title in hidden:
            sheet.sheet_state = "hidden"
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def sheet_titles(data: bytes) -> list[str]:
    return [s.title for s in openpyxl.load_workbook(BytesIO(data)).worksheets]


@pytest.fixture
def limits(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "max_pages", 50)
    monkeypatch.setattr(settings, "max_sheets", 3)
    monkeypatch.setattr(settings, "large_document_mode", "partial")


# --- PDF ----------------------------------------------------------------------------


def test_small_pdf_is_read_in_full(limits: None) -> None:
    prepared = prepare("week.pdf", pdf(3))
    assert prepared.page_range is None
    assert not prepared.limit.content_left_out
    assert too_large_flag(prepared.limit.as_dict()) is None


def test_100_page_pdf_reads_the_first_50_and_flags_it(limits: None) -> None:
    prepared = prepare("big.pdf", pdf(100))
    assert prepared.page_range == (1, 50)
    assert prepared.limit.content_left_out
    flag = too_large_flag(prepared.limit.as_dict())
    assert flag is not None and flag["code"] == "document_too_large"
    assert "only pages 1-50 of 100 were read" in flag["message"]


def test_refuse_mode_says_split_the_file(limits: None, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "large_document_mode", "refuse")
    with pytest.raises(DocumentTooLargeError, match="has 100 pages; the limit is 50. Split"):
        prepare("big.pdf", pdf(100))


def test_damaged_pdf_is_left_to_docling(limits: None) -> None:
    prepared = prepare("broken.pdf", b"%PDF-1.7 not really")
    assert prepared.page_range is None and prepared.data == b"%PDF-1.7 not really"


# --- Excel --------------------------------------------------------------------------


def test_hidden_and_empty_sheets_are_never_read(limits: None) -> None:
    data = workbook(
        {"Week 1": [["Name", "Mon"], ["Grace", 8]], "Empty": [], "Rates": [["x", 1]]},
        hidden=("Rates",),
    )
    prepared = prepare("hours.xlsx", data)
    assert sheet_titles(prepared.data) == ["Week 1"]
    assert prepared.limit.sheets_skipped_empty == ["Empty"]
    assert prepared.limit.sheets_skipped_hidden == ["Rates"]
    assert not prepared.limit.content_left_out  # nothing with data was left out
    assert too_large_flag(prepared.limit.as_dict()) is None


def test_too_many_sheets_keeps_the_first_n_and_flags_it(limits: None) -> None:
    data = workbook({f"Week {n}": [["Grace", 8]] for n in range(1, 6)})
    prepared = prepare("hours.xlsx", data)
    assert sheet_titles(prepared.data) == ["Week 1", "Week 2", "Week 3"]
    flag = too_large_flag(prepared.limit.as_dict())
    assert flag is not None and "2 sheet(s) were not read (Week 4, Week 5)" in flag["message"]


def test_a_normal_workbook_is_not_rewritten(limits: None) -> None:
    data = workbook({"Week 1": [["Grace", 8]]})
    assert prepare("hours.xlsx", data).data == data


def test_other_formats_pass_through(limits: None) -> None:
    prepared = prepare("photo.jpg", b"\xff\xd8 image")
    assert prepared.data == b"\xff\xd8 image" and prepared.page_range is None


def test_settings() -> None:
    defaults = Settings(_env_file=None)
    assert (defaults.max_pages, defaults.max_sheets) == (50, 30)
    assert defaults.large_document_mode == "partial"
    assert defaults.llm_max_output_tokens == 16_000


# --- errors in words people understand ------------------------------------------------


class BadRequestError(Exception):
    """Looks like the provider's error (module and class name are what explain() checks)."""


BadRequestError.__module__ = "openai"


@pytest.mark.parametrize(
    "text",
    [
        "Error code: 400 - {'error': {'code': 'context_length_exceeded'}}",
        "This model's maximum context length is 131072 tokens",
        "prompt is too long: 250000 tokens > 200000 maximum",
    ],
)
def test_too_long_for_the_model_is_not_blamed_on_the_api_key(text: str) -> None:
    failure = explain(BadRequestError(text))
    assert failure.message == "The document is too long for the AI model."
    assert failure.temporary is False
    assert "API key" not in failure.message


def test_a_real_setup_error_still_says_api_key() -> None:
    failure = explain(BadRequestError("Invalid model name 'gpt-9'"))
    assert "API key or model name" in failure.message


def test_a_cut_off_answer_is_permanent_with_advice() -> None:
    failure = explain(AnswerTooLongError("cut off twice"))
    assert failure.temporary is False
    assert "too many people or days" in failure.message
    assert "LLM_MAX_OUTPUT_TOKENS" in failure.advice


def test_cut_off_twice_stops_after_one_retry() -> None:
    pytest.importorskip("langchain_core")
    from src.control.agents.structured_output import AnswerCutOffError, invoke_structured
    from src.schemas.extraction import ExtractionResult

    calls = 0

    class CutOffModel:
        async def ainvoke(self, messages: Any) -> Any:
            nonlocal calls
            calls += 1
            return SimpleNamespace(
                content='{"timesheets": [', response_metadata={"finish_reason": "length"}
            )

    with pytest.raises(AnswerCutOffError):
        asyncio.run(
            invoke_structured(
                cast(Any, CutOffModel()),
                ExtractionResult,
                system="s",
                user="u",
                retries=3,
            )
        )
    assert calls == 2  # one retry asking for a shorter answer, then stop (not 4 calls)
