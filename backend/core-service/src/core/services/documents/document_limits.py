"""Size checks BEFORE a document is read: pages of a PDF, sheets of an Excel file.

Reading is the slow part (Docling, OCR for scans), and the AI only sees the first
LLM_MAX_INPUT_CHARS characters anyway. So instead of converting 100 pages and throwing
most of the text away, we look first (cheap: page / sheet count) and then:

  PDF   > MAX_PAGES  -> partial: read pages 1..MAX_PAGES      | refuse: "split the file"
  XLSX               -> hidden and empty sheets are dropped (never read)
        > MAX_SHEETS -> partial: keep the first MAX_SHEETS     | refuse: "split the file"

LARGE_DOCUMENT_MODE chooses partial (default) or refuse. What was left out is returned
as a ReadLimit, stored on the document, and becomes a document_too_large flag on every
timesheet from it, so a reviewer knows pages or sheets were not read.

Uses pypdfium2 and openpyxl, which come with Docling (no new dependency)."""

from dataclasses import asdict, dataclass, field
from io import BytesIO
from typing import Any

from src.config.settings import settings
from src.utils.file_checks import zip_bomb_reason


class DocumentTooLargeError(Exception):
    """LARGE_DOCUMENT_MODE=refuse and the document is over the limit."""


@dataclass(slots=True)
class ReadLimit:
    """What was (not) read. Empty lists / None = nothing was left out."""

    pages_total: int | None = None
    pages_read: int | None = None
    sheets_total: int | None = None
    sheets_read: list[str] = field(default_factory=list)
    sheets_skipped_over_limit: list[str] = field(default_factory=list)
    sheets_skipped_empty: list[str] = field(default_factory=list)
    sheets_skipped_hidden: list[str] = field(default_factory=list)

    @property
    def content_left_out(self) -> bool:
        """True when real content (not just empty / hidden sheets) was not read."""
        pages_cut = (
            self.pages_total is not None
            and self.pages_read is not None
            and self.pages_read < self.pages_total
        )
        return pages_cut or bool(self.sheets_skipped_over_limit)

    def as_dict(self) -> dict[str, Any]:
        return {k: v for k, v in asdict(self).items() if v not in (None, [])}


@dataclass(slots=True)
class Prepared:
    data: bytes  # possibly rewritten (Excel without the skipped sheets)
    page_range: tuple[int, int] | None  # PDF pages to read (1-based, inclusive)
    limit: ReadLimit


def _too_large(what: str, total: int, limit: int, unit: str) -> DocumentTooLargeError:
    return DocumentTooLargeError(
        f"'{what}' has {total} {unit}; the limit is {limit}. Split it into smaller "
        "files and send them again."
    )


def pdf_page_count(data: bytes) -> int:
    import pypdfium2

    document = pypdfium2.PdfDocument(data)
    try:
        return len(document)
    finally:
        document.close()


def _prepare_pdf(filename: str, data: bytes) -> Prepared:
    try:
        pages = pdf_page_count(data)
    except Exception:  # damaged / encrypted: let Docling report it the usual way
        return Prepared(data, None, ReadLimit())
    limit = settings.max_pages
    if pages <= limit:
        return Prepared(data, None, ReadLimit(pages_total=pages, pages_read=pages))
    if settings.large_document_mode == "refuse":
        raise _too_large(filename, pages, limit, "pages")
    return Prepared(data, (1, limit), ReadLimit(pages_total=pages, pages_read=limit))


def _sheet_is_empty(sheet: Any) -> bool:
    for row in sheet.iter_rows(values_only=True):
        if any(cell not in (None, "") for cell in row):
            return False
    return True


def _prepare_xlsx(filename: str, data: bytes) -> Prepared:
    try:
        import openpyxl

        book = openpyxl.load_workbook(BytesIO(data), data_only=True)
    except Exception:  # not a real workbook: let Docling report it the usual way
        return Prepared(data, None, ReadLimit())

    limit = ReadLimit(sheets_total=len(book.worksheets))
    keep: list[Any] = []
    for sheet in book.worksheets:
        if sheet.sheet_state != "visible":
            limit.sheets_skipped_hidden.append(sheet.title)
        elif _sheet_is_empty(sheet):
            limit.sheets_skipped_empty.append(sheet.title)
        else:
            keep.append(sheet)

    if len(keep) > settings.max_sheets:
        if settings.large_document_mode == "refuse":
            raise _too_large(filename, len(keep), settings.max_sheets, "sheets with data")
        limit.sheets_skipped_over_limit = [s.title for s in keep[settings.max_sheets :]]
        keep = keep[: settings.max_sheets]
    limit.sheets_read = [s.title for s in keep]

    if len(keep) == len(book.worksheets):
        return Prepared(data, None, limit)  # nothing to remove: read the original bytes
    if not keep:
        return Prepared(data, None, limit)  # all empty: Docling says "no readable text"
    for sheet in list(book.worksheets):
        if sheet not in keep:
            book.remove(sheet)
    book.active = 0
    out = BytesIO()
    book.save(out)
    return Prepared(out.getvalue(), None, limit)


def prepare(filename: str, data: bytes) -> Prepared:
    """Check the size and cut the document down if needed. Raises DocumentTooLargeError
    in refuse mode. Other formats (images, DOCX) pass through unchanged."""
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension in ("docx", "xlsx"):
        # Second line of defence (intake already checked): never hand a zip bomb to a
        # reader, even for a document that reached the worker another way.
        reason = zip_bomb_reason(
            filename,
            data,
            max_uncompressed_mb=settings.max_uncompressed_mb,
            max_ratio=settings.max_compression_ratio,
        )
        if reason is not None:
            raise DocumentTooLargeError(reason)
    if extension == "pdf":
        return _prepare_pdf(filename, data)
    if extension == "xlsx":
        return _prepare_xlsx(filename, data)
    return Prepared(data, None, ReadLimit())


def too_large_flag(limit: dict[str, Any] | None) -> dict[str, Any] | None:
    """Pure: the document_too_large flag for a stored ReadLimit, or None."""
    if not limit:
        return None
    parts: list[str] = []
    total, read = limit.get("pages_total"), limit.get("pages_read")
    if total and read and read < total:
        parts.append(f"only pages 1-{read} of {total} were read")
    over = limit.get("sheets_skipped_over_limit") or []
    if over:
        shown = ", ".join(over[:5]) + ("…" if len(over) > 5 else "")
        parts.append(f"{len(over)} sheet(s) were not read ({shown})")
    if not parts:
        return None
    return {
        "code": "document_too_large",
        "message": "The document is too large: "
        + "; ".join(parts)
        + ". Hours on those pages or sheets are missing. Split the file and send it again.",
        **{k: v for k, v in limit.items() if k in ("pages_total", "pages_read")},
    }
