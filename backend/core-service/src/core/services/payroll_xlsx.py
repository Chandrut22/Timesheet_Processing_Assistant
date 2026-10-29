"""The payroll sheet as an Excel workbook (.xlsx), built from the stored CSV.

The CSV saved with each export batch stays the record of what was exported. The
workbook is made from it on download, so both formats always hold the same rows,
also for exports made before .xlsx existed.

What the workbook adds: real numbers and dates (not text), number formats, a frozen
header with filters, column widths and a bold totals row. Pure function (no
database), so the format is unit-tested."""

import csv
import io
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.core.services.payroll_csv import _FORMULA_START, CSV_COLUMNS

SHEET_TITLE = "Payroll"
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HOUR_COLUMNS = frozenset(
    {
        "regular_hours",
        "overtime_hours",
        "double_time_hours",
        "holiday_hours",
        "total_hours",
    }
)
RATE_COLUMNS = frozenset(
    {"hourly_rate", "overtime_rate", "double_time_rate", "holiday_rate"}
)
AMOUNT_COLUMNS = frozenset(
    {
        "regular_amount",
        "overtime_amount",
        "double_time_amount",
        "holiday_amount",
        "gross_amount",
    }
)
DATE_COLUMNS = frozenset({"week_start", "week_end"})
DATETIME_COLUMNS = frozenset({"approved_at"})
TOTAL_COLUMNS = HOUR_COLUMNS | AMOUNT_COLUMNS  # rates are not added up

HOURS_FORMAT = "0.00"
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"  # UTC, as in the CSV

WIDTHS = {
    "company": 28,
    "employee_name": 26,
    "employee_code": 14,
    "approved_by": 20,
    "approved_at": 17,
    "timesheet_id": 38,
}


def _text(value: str) -> str:
    """Undo the CSV's formula guard ("'=X" -> "=X"): in the workbook the cell is
    written as text, so Excel never runs it."""
    if len(value) > 1 and value[0] == "'" and value[1:].startswith(_FORMULA_START):
        return value[1:]
    return value


def _number(value: str) -> Decimal | None:
    if value == "":
        return None
    try:
        return Decimal(value)
    except InvalidOperation:
        return None


def _cell_value(column: str, raw: str) -> object:
    if column in HOUR_COLUMNS | RATE_COLUMNS | AMOUNT_COLUMNS:
        number = _number(raw)
        return float(number) if number is not None else None
    if column in DATE_COLUMNS and raw:
        return date.fromisoformat(raw)
    if column in DATETIME_COLUMNS and raw:
        moment = datetime.fromisoformat(raw)
        if moment.tzinfo is not None:  # Excel has no time zones: write UTC
            moment = moment.astimezone(UTC).replace(tzinfo=None)
        return moment
    return _text(raw) if raw else None


def _number_format(column: str) -> str | None:
    if column in HOUR_COLUMNS:
        return HOURS_FORMAT
    if column in RATE_COLUMNS | AMOUNT_COLUMNS:
        return MONEY_FORMAT
    if column in DATE_COLUMNS:
        return DATE_FORMAT
    if column in DATETIME_COLUMNS:
        return DATETIME_FORMAT
    return None


def build_xlsx(csv_content: str, *, title: str = "Payroll export") -> bytes:
    """The stored payroll CSV -> an .xlsx file (bytes), one row per timesheet plus a
    totals row."""
    reader = csv.reader(io.StringIO(csv_content.lstrip("﻿")))
    header = next(reader, None) or list(CSV_COLUMNS)
    data = [line for line in reader if line]

    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_TITLE
    book.properties.title = title[:255]

    bold = Font(bold=True)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = bold

    sums: dict[str, Decimal] = {c: Decimal("0") for c in header if c in TOTAL_COLUMNS}
    for line in data:
        values = []
        for column, raw in zip(header, line, strict=False):
            values.append(_cell_value(column, raw))
            if column in sums:
                sums[column] += _number(raw) or Decimal("0")
        sheet.append(values)
        row = sheet.max_row
        for index, column in enumerate(header, start=1):
            cell = sheet.cell(row=row, column=index)
            number_format = _number_format(column)
            if number_format and cell.value is not None:
                cell.number_format = number_format
            elif isinstance(cell.value, str):
                cell.data_type = "s"  # text stays text, even if it starts with "="

    last_data_row = sheet.max_row
    if data:
        totals_row = last_data_row + 1
        for index, column in enumerate(header, start=1):
            cell = sheet.cell(row=totals_row, column=index)
            if index == 1:
                cell.value = "Total"
            elif column in sums:
                cell.value = float(sums[column])
                cell.number_format = _number_format(column) or HOURS_FORMAT
            cell.font = bold

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{last_data_row}"
    for index, column in enumerate(header, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS.get(
            column, max(12, len(column) + 2)
        )

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def xlsx_filename(csv_filename: str) -> str:
    """payroll_acme_2026-09-14_1a2b3c4d.csv -> payroll_acme_2026-09-14_1a2b3c4d.xlsx"""
    stem = csv_filename[:-4] if csv_filename.lower().endswith(".csv") else csv_filename
    return f"{stem}.xlsx"
